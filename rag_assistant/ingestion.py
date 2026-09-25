"""Document loading, validation, normalization, and chunking."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Protocol

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx"}
_WHITESPACE_RE = re.compile(r"[ \t]+")
_MULTI_NEWLINE_RE = re.compile(r"\n{3,}")


class UploadedFileLike(Protocol):
    name: str

    def getvalue(self) -> bytes: ...


@dataclass(frozen=True)
class FilePayload:
    """A framework-independent representation of an uploaded file."""

    name: str
    data: bytes
    content_type: str | None = None


@dataclass
class IngestionResult:
    documents: list[Document] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    files_processed: int = 0
    pages_processed: int = 0
    characters: int = 0


def payload_from_upload(uploaded_file: UploadedFileLike) -> FilePayload:
    """Convert a Streamlit UploadedFile into a FilePayload."""

    return FilePayload(
        name=str(getattr(uploaded_file, "name", "uploaded-document")),
        data=bytes(uploaded_file.getvalue()),
        content_type=getattr(uploaded_file, "type", None),
    )


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = _WHITESPACE_RE.sub(" ", text)
    return _MULTI_NEWLINE_RE.sub("\n\n", text).strip()


def _base_metadata(payload: FilePayload, extension: str) -> dict[str, str | int]:
    return {
        "source": Path(payload.name).name,
        "file_type": extension.lstrip("."),
        "content_hash": hashlib.sha256(payload.data).hexdigest()[:16],
    }


def _extract_pdf(payload: FilePayload, metadata: dict[str, str | int]) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(payload.data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("the PDF is password-protected") from exc

    documents: list[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = _normalize_text(page.extract_text() or "")
        if not text:
            continue
        page_metadata = dict(metadata)
        page_metadata.update({"page": page_number, "page_count": len(reader.pages)})
        documents.append(Document(page_content=text, metadata=page_metadata))

    if not documents:
        raise ValueError("the PDF contains no extractable text (it may be scanned)")
    return documents


def _extract_docx(payload: FilePayload, metadata: dict[str, str | int]) -> list[Document]:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(payload.data))
    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    text = _normalize_text("\n".join(parts))
    if not text:
        raise ValueError("the DOCX contains no extractable text")
    docx_metadata = dict(metadata)
    docx_metadata.update({"page": 1, "page_count": 1})
    return [Document(page_content=text, metadata=docx_metadata)]


def _extract_text(payload: FilePayload, metadata: dict[str, str | int]) -> list[Document]:
    text = _normalize_text(payload.data.decode("utf-8-sig", errors="replace"))
    if not text:
        raise ValueError("the file contains no readable text")
    text_metadata = dict(metadata)
    text_metadata.update({"page": 1, "page_count": 1})
    return [Document(page_content=text, metadata=text_metadata)]


def _extract_documents(payload: FilePayload, extension: str) -> list[Document]:
    metadata = _base_metadata(payload, extension)
    if extension == ".pdf":
        return _extract_pdf(payload, metadata)
    if extension == ".docx":
        return _extract_docx(payload, metadata)
    if extension in {".txt", ".md", ".markdown"}:
        return _extract_text(payload, metadata)
    raise ValueError(f"unsupported file type '{extension}'")


def load_file_payloads(
    payloads: Iterable[FilePayload],
    *,
    max_file_size_mb: int = 10,
    max_files: int = 10,
) -> IngestionResult:
    """Load and validate files, returning page-aware LangChain Documents."""

    result = IngestionResult()
    selected_payloads = list(payloads)
    if len(selected_payloads) > max_files:
        skipped = len(selected_payloads) - max_files
        result.errors.append(
            f"Only the first {max_files} files were processed; {skipped} were skipped."
        )
        selected_payloads = selected_payloads[:max_files]

    max_bytes = max_file_size_mb * 1024 * 1024
    seen_hashes: set[str] = set()

    for payload in selected_payloads:
        name = Path(payload.name).name
        extension = Path(name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            result.errors.append(f"{name}: unsupported file type '{extension or 'unknown'}'.")
            continue
        if not payload.data:
            result.errors.append(f"{name}: file is empty.")
            continue
        if len(payload.data) > max_bytes:
            result.errors.append(
                f"{name}: exceeds the {max_file_size_mb} MB file-size limit."
            )
            continue

        content_hash = hashlib.sha256(payload.data).hexdigest()[:16]
        if content_hash in seen_hashes:
            result.errors.append(f"{name}: duplicate of an earlier upload was skipped.")
            continue

        try:
            documents = _extract_documents(payload, extension)
        except Exception as exc:
            result.errors.append(f"{name}: {exc}")
            continue

        seen_hashes.add(content_hash)
        result.documents.extend(documents)
        result.files_processed += 1
        result.pages_processed += len(documents)
        result.characters += sum(len(document.page_content) for document in documents)

    return result


def load_uploaded_files(
    uploaded_files: Iterable[UploadedFileLike],
    *,
    max_file_size_mb: int = 10,
    max_files: int = 10,
) -> IngestionResult:
    """Load Streamlit UploadedFile objects."""

    return load_file_payloads(
        (payload_from_upload(uploaded_file) for uploaded_file in uploaded_files),
        max_file_size_mb=max_file_size_mb,
        max_files=max_files,
    )


def load_file_payloads_from_paths(
    paths: Iterable[Path], *, max_file_size_mb: int = 10, max_files: int = 10
) -> IngestionResult:
    """Load files from disk, primarily for local demos and evaluation."""

    payloads: list[FilePayload] = []
    for path in sorted(Path(path) for path in paths):
        if path.is_file():
            payloads.append(FilePayload(name=path.name, data=path.read_bytes()))
    return load_file_payloads(
        payloads, max_file_size_mb=max_file_size_mb, max_files=max_files
    )


def chunk_documents(
    documents: Iterable[Document],
    *,
    chunk_size: int = 1_000,
    chunk_overlap: int = 150,
    max_chunks: int = 6_000,
) -> list[Document]:
    """Split documents while preserving source, page, and chunk metadata."""

    if chunk_size < 256:
        raise ValueError("chunk_size must be at least 256 characters")
    if chunk_overlap < 0 or chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be non-negative and smaller than chunk_size")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks: list[Document] = []
    for source_document in documents:
        splits = splitter.split_documents([source_document])
        for chunk_index, split in enumerate(splits):
            content = split.page_content.strip()
            if not content:
                continue
            metadata = dict(split.metadata)
            metadata.update(
                {
                    "chunk_index": chunk_index,
                    "chunk_id": (
                        f"{metadata.get('source', 'document')}:"
                        f"{metadata.get('page', 1)}:{chunk_index}"
                    ),
                }
            )
            chunks.append(Document(page_content=content, metadata=metadata))
            if len(chunks) >= max_chunks:
                return chunks

    return chunks


def ingest_and_chunk(
    payloads: Iterable[FilePayload], settings: object
) -> tuple[IngestionResult, list[Document]]:
    """Convenience function used by scripts and tests."""

    result = load_file_payloads(
        payloads,
        max_file_size_mb=getattr(settings, "max_file_size_mb", 10),
        max_files=getattr(settings, "max_files", 10),
    )
    chunks = chunk_documents(
        result.documents,
        chunk_size=getattr(settings, "chunk_size", 1_000),
        chunk_overlap=getattr(settings, "chunk_overlap", 150),
        max_chunks=getattr(settings, "max_chunks", 6_000),
    )
    return result, chunks
